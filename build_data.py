#!/usr/bin/env python3
"""
build_data.py - TSET-AHIF Project Plan Progress Dashboard
=========================================================

Reads the "Dashboard_Data" tab of the project-plan workbook and writes
data/dashboard.json, which index.html loads in the browser.

Usage (from this folder):
    python build_data.py
    python build_data.py "C:/path/to/TSET_Project_Plan_Tracking_Data.xlsx"

If no path is given, the script looks for the workbook in this order:
    1. ./data/TSET_Project_Plan_Tracking_Data.xlsx
    2. ../Data/TSET_Project_Plan_Tracking_Data.xlsx   (the shared "Data" folder)
    3. ../../Data/TSET_Project_Plan_Tracking_Data.xlsx (when run from inside Claude_Output)

IMPORTANT: the Dashboard_Data tab is filled by formulas that pull from the
"Project Plan" tab. This script reads the values Excel saved with the file,
so always open the workbook in Excel and SAVE it after editing before you
run this script.
"""

import json
import re
import sys
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required. Install it with:  pip install openpyxl")

HERE = Path(__file__).resolve().parent
WORKBOOK_NAME = "TSET_Project_Plan_Tracking_Data.xlsx"
SHEET_NAME = "Dashboard_Data"
OUT_PATH = HERE / "data" / "dashboard.json"

DASHBOARD_TITLE = "TSET-AHIF Project Plan Progress Dashboard"

# Allowed status values (in progress order). Anything else is flagged.
STATUSES = ["Not Started", "In Progress", "Complete"]
CATEGORIES = ["Construction", "Administrative"]

# Header text in the sheet -> internal key. Matching ignores case/extra spaces.
COLUMNS = {
    "year": "year",
    "project_type": "category",
    "project milestones": "milestone",
    "task assigned to": "assigned_to",
    "target start date": "target_start",
    "target end date": "target_end",
    "actual start date": "actual_start",
    "actual end date": "actual_end",
    "task status": "status",
    "notes": "notes",
}
DATE_KEYS = ["target_start", "target_end", "actual_start", "actual_end"]
EMPTY_TOKENS = {"", "n/a", "na", "none", "tbd", "-", "0"}


def clean_text(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)) and v == 0:
        return ""          # formula spill writes 0 for blank source cells
    return re.sub(r"\s+", " ", str(v)).strip()


def parse_date(v):
    """Return (iso_string_or_None, problem_text_or_None)."""
    if v is None:
        return None, None
    if isinstance(v, datetime):
        return v.date().isoformat(), None
    if isinstance(v, date):
        return v.isoformat(), None
    s = clean_text(v)
    if s.lower() in EMPTY_TOKENS:
        return None, None
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat(), None
        except ValueError:
            pass
    return None, s


def normalize_status(v):
    s = clean_text(v).lower()
    for st in STATUSES:
        if s == st.lower():
            return st
    if s in ("", "not started", "notstarted"):
        return "Not Started"
    if "progress" in s:
        return "In Progress"
    if s.startswith("complete") or s in ("done", "completed"):
        return "Complete"
    return None


def normalize_category(v):
    s = clean_text(v)
    for c in CATEGORIES:
        if s.lower() == c.lower():
            return c
    return s or "Uncategorized"


def find_workbook():
    if len(sys.argv) > 1:
        p = Path(sys.argv[1]).expanduser()
        if not p.exists():
            sys.exit(f"Workbook not found: {p}")
        return p
    for p in (HERE / "data" / WORKBOOK_NAME, HERE.parent / "Data" / WORKBOOK_NAME,
              HERE.parent.parent / "Data" / WORKBOOK_NAME):
        if p.exists():
            return p
    sys.exit(
        f"Could not find {WORKBOOK_NAME}.\n"
        f"Put it in ./data/ or ../Data/, or pass its path:  python build_data.py <path>"
    )


def main():
    wb_path = find_workbook()
    print(f"Reading: {wb_path}")
    wb = openpyxl.load_workbook(wb_path, data_only=True)  # data_only -> Excel's saved values
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"Tab '{SHEET_NAME}' not found. Tabs present: {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        sys.exit("Dashboard_Data tab is empty.")

    # Map header positions
    header = [clean_text(h).lower() for h in rows[0]]
    idx = {}
    for i, h in enumerate(header):
        if h in COLUMNS:
            idx[COLUMNS[h]] = i
    missing = [k for k in COLUMNS.values() if k not in idx]
    if missing:
        sys.exit(f"Missing expected column(s) in {SHEET_NAME}: {missing}\nFound headers: {rows[0]}")

    issues = []          # shown in the dashboard's Data-quality notes
    milestones = []
    year_labels = {}     # year number -> list of calendar years seen
    empty_formula_rows = 0

    for r_i, row in enumerate(rows[1:], start=2):
        get = lambda k: row[idx[k]] if idx[k] < len(row) else None
        name = clean_text(get("milestone"))
        year_raw = clean_text(get("year"))
        if not name and not year_raw:
            continue
        if not name:
            empty_formula_rows += 1
            continue

        m = re.search(r"year\s*(\d+)", year_raw, re.I)
        year_num = int(m.group(1)) if m else None
        cal = re.match(r"\s*(\d{4})", year_raw)
        if year_num is None:
            issues.append(f"Row {r_i}: Year value \"{year_raw}\" doesn't contain \"Year N\" - listed under \"Unassigned year\".")
        else:
            year_labels.setdefault(year_num, [])
            if cal:
                year_labels[year_num].append(int(cal.group(1)))

        rec = OrderedDict()
        rec["id"] = f"m{r_i}"
        rec["row"] = r_i
        rec["year"] = year_num
        rec["year_raw"] = year_raw
        rec["category"] = normalize_category(get("category"))
        rec["milestone"] = name
        rec["assigned_to"] = clean_text(get("assigned_to"))

        for k in DATE_KEYS:
            iso, bad = parse_date(get(k))
            rec[k] = iso
            if bad:
                rec[k + "_text"] = bad
                issues.append(
                    f"\"{name}\": {k.replace('_', ' ')} \"{bad}\" isn't a valid date - shown as text."
                )

        st = normalize_status(get("status"))
        if st is None:
            issues.append(f"\"{name}\": status \"{clean_text(get('status'))}\" not recognized - treated as Not Started.")
            st = "Not Started"
        rec["status"] = st
        rec["notes"] = clean_text(get("notes"))
        milestones.append(rec)

    # Year labels: "Year 1 (2026)" using the earliest calendar year typed for that year number.
    years = []
    for n in sorted(year_labels):
        cals = sorted(set(year_labels[n]))
        cal = cals[0] if cals else None
        if len(cals) > 1:
            issues.append(
                f"Year {n}: the Year column mixes calendar years ({', '.join(map(str, cals))}); "
                f"grouped together as Year {n} ({cal})."
            )
        years.append({"num": n, "label": f"Year {n}", "calendar": cal})
    if any(m["year"] is None for m in milestones):
        years.append({"num": None, "label": "Unassigned year", "calendar": None})

    # Summary counts
    def counts(items):
        c = {s: 0 for s in STATUSES}
        for it in items:
            c[it["status"]] += 1
        return c

    summary = {
        "total": len(milestones),
        "by_status": counts(milestones),
        "by_year": {str(y["num"]): counts([m for m in milestones if m["year"] == y["num"]]) for y in years},
        "by_category": {c: counts([m for m in milestones if m["category"] == c]) for c in CATEGORIES},
    }

    out = OrderedDict()
    out["title"] = DASHBOARD_TITLE
    out["generated_at"] = datetime.now().isoformat(timespec="minutes")
    out["source_file"] = wb_path.name
    out["source_modified"] = datetime.fromtimestamp(wb_path.stat().st_mtime).isoformat(timespec="minutes")
    out["statuses"] = STATUSES
    out["categories"] = CATEGORIES
    out["years"] = years
    out["summary"] = summary
    out["milestones"] = milestones
    out["issues"] = issues

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"Wrote:   {OUT_PATH}")
    print(f"Milestones: {len(milestones)}  |  " +
          "  ".join(f"{s}: {summary['by_status'][s]}" for s in STATUSES))
    if empty_formula_rows:
        print(f"\nWARNING: {empty_formula_rows} row(s) had a Year but no milestone text.")
        print("  If the whole year block is missing, open the workbook in Excel, save it, and rerun.")
    if issues:
        print("\nData-quality notes (also shown on the dashboard):")
        for i in issues:
            print("  - " + i)


if __name__ == "__main__":
    main()
